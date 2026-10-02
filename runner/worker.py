"""Bounded worker protocol. This program is never run on the application host."""
import base64
import hashlib
import json
import math
import mimetypes
import os
import re
import selectors
import signal
import socket
import http.server
import queue
import uuid
import stat
import subprocess
import sys
import threading
import time
import zipfile
from pathlib import Path

if os.getpid() != 1 or os.getuid() == 0:
    raise SystemExit('The worker requires a non-root private container PID namespace (PID 1).')

ROOT = Path('/workspace')
FILE_LIMIT = 16 * 1024 * 1024
TOTAL_LIMIT = 32 * 1024 * 1024
OUTPUT_LIMIT = 2 * 1024 * 1024
CHUNK = 192 * 1024
bridge = None
input_stream = None
input_hash = None
input_size = 0
input_expected = None
sealed = False
requests = queue.Queue(maxsize=2)
registry_replies = {}
output_lock = threading.Lock()


def emit(value):
    with output_lock:
        sys.stdout.write(json.dumps(value, ensure_ascii=True) + '\n')
        sys.stdout.flush()


def receive():
    for line in iter(lambda: sys.stdin.buffer.readline(1024 * 1024 + 1), b''):
        if len(line) > 1024 * 1024:
            break
        try:
            message = json.loads(line)
            if message.get('op') == 'registry-response':
                reply = registry_replies.get(message.get('registryId'))
                if reply:
                    reply.put(message, timeout=5)
            else:
                requests.put(message)
        except (ValueError, queue.Full):
            break
    requests.put(None)


def relative(value):
    if not isinstance(value, str) or not value or len(value) > 240 or re.search(r'[\x00-\x1f\x7f\\]', value):
        raise ValueError('Invalid relative path.')
    parts = value.split('/')
    if len(parts) > 12 or any(part in ('', '.', '..') for part in parts):
        raise ValueError('Invalid relative path.')
    return ROOT / value


def child_env(packages=False):
    return {'PATH': '/deps/node/node_modules/.bin:/usr/local/bin:/usr/bin:/bin' if packages else '/usr/local/bin:/usr/bin:/bin',
            'HOME': '/tmp/home', 'LANG': 'C.UTF-8', 'TMPDIR': '/scratch',
            'MPLBACKEND': 'Agg', 'MPLCONFIGDIR': '/tmp/matplotlib', 'PYTHONDONTWRITEBYTECODE': '1',
            'PYTHONPATH': '/deps/python' if packages else '', 'NODE_PATH': '/deps/node/node_modules' if packages else '',
            'PIP_CONFIG_FILE': '/dev/null', 'PIP_DISABLE_PIP_VERSION_CHECK': '1',
            'NPM_CONFIG_USERCONFIG': '/tmp/empty-user-npmrc', 'NPM_CONFIG_GLOBALCONFIG': '/tmp/empty-global-npmrc'}


def kill_children():
    # The supervisor is PID 1 in a private PID namespace. Also stop detached children.
    for _ in range(3):
        for entry in Path('/proc').iterdir():
            if entry.name.isdigit() and int(entry.name) > 1:
                try:
                    os.kill(int(entry.name), signal.SIGKILL)
                except ProcessLookupError:
                    pass
        while True:
            try:
                pid, _status = os.waitpid(-1, os.WNOHANG)
                if not pid:
                    break
            except ChildProcessError:
                break
        time.sleep(0.02)


def command(argv, timeout, cwd='/workspace', packages=False, cap=OUTPUT_LIMIT):
    proc = subprocess.Popen(argv, cwd=cwd, env=child_env(packages), stdin=subprocess.DEVNULL,
                            stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True)
    poller = selectors.DefaultSelector()
    poller.register(proc.stdout, selectors.EVENT_READ, 'stdout')
    poller.register(proc.stderr, selectors.EVENT_READ, 'stderr')
    chunks = {'stdout': [], 'stderr': []}
    size = 0
    deadline = time.monotonic() + timeout
    failure = None
    try:
        while poller.get_map():
            if time.monotonic() >= deadline:
                failure = 'Execution timed out.'
                break
            for key, _mask in poller.select(0.1):
                data = os.read(key.fileobj.fileno(), 65536)
                if not data:
                    poller.unregister(key.fileobj)
                    continue
                remaining = max(0, cap - size)
                chunks[key.data].append(data[:remaining])
                size += len(data)
                if size > cap:
                    failure = 'Output limit exceeded.'
                    break
            if failure:
                break
        if failure:
            os.killpg(proc.pid, signal.SIGKILL)
        try:
            status = proc.wait(timeout=max(0.1, deadline - time.monotonic()))
        except subprocess.TimeoutExpired:
            failure = 'Execution timed out.'
            os.killpg(proc.pid, signal.SIGKILL)
            status = proc.wait(timeout=2)
    finally:
        poller.close()
        proc.stdout.close()
        proc.stderr.close()
        kill_children()
    return {'stdout': b''.join(chunks['stdout']).decode('utf-8', 'replace'),
            'stderr': b''.join(chunks['stderr']).decode('utf-8', 'replace'), 'exitCode': status,
            'status': 'timed_out' if failure == 'Execution timed out.' else 'error' if failure or status else 'complete',
            'error': failure or (f'Process exited with status {status}.' if status else None)}


class Bridge(http.server.ThreadingHTTPServer):
    allow_reuse_address = True
    daemon_threads = True

    def __init__(self):
        self.connections = set()
        self.closed = False
        super().__init__(('127.0.0.1', 43123), Relay)
        self.thread = threading.Thread(target=self.serve_forever, daemon=True)
        self.thread.start()

    def close(self):
        self.closed = True
        self.shutdown()
        for connection in list(self.connections):
            try:
                connection.shutdown(socket.SHUT_RDWR)
            except OSError:
                pass
            connection.close()
        self.server_close()
        self.thread.join(timeout=2)


class Relay(http.server.BaseHTTPRequestHandler):
    protocol_version = 'HTTP/1.1'

    def log_message(self, _format, *_args):
        pass

    def do_HEAD(self):
        self.do_GET()

    def do_GET(self):
        if self.server.closed or len(registry_replies) >= 16 or len(self.path) > 2048:
            self.send_error(503, 'Registry access unavailable.')
            return
        request_id = str(uuid.uuid4())
        replies = queue.Queue(maxsize=4)
        registry_replies[request_id] = replies
        self.server.connections.add(self.connection)
        self.connection.settimeout(25)
        try:
            emit({'event': 'registry', 'registryId': request_id, 'method': self.command, 'path': self.path})
            start = replies.get(timeout=25)
            if start.get('stage') != 'start' or not 0 <= start['size'] <= 64 * 1024 * 1024:
                raise ValueError('Invalid registry response.')
            self.send_response(start['status'])
            self.send_header('Content-Type', start['type'])
            self.send_header('Content-Length', str(start['size']))
            self.end_headers()
            size = 0
            while not self.server.closed:
                part = replies.get(timeout=25)
                if part.get('stage') == 'end':
                    if size != start['size']:
                        raise ValueError('Incomplete registry response.')
                    return
                data = base64.b64decode(part['data'], validate=True)
                size += len(data)
                if len(data) > CHUNK or size > start['size']:
                    raise ValueError('Registry chunk exceeds its bound.')
                self.wfile.write(data)
                self.wfile.flush()
        except (OSError, ValueError, KeyError, queue.Empty):
            self.close_connection = True
        finally:
            registry_replies.pop(request_id, None)
            self.server.connections.discard(self.connection)


def install(packages):
    global bridge
    if sealed or list(ROOT.iterdir()):
        raise ValueError('Packages must be installed in clean scratch before files arrive.')
    if bridge is not None:
        raise ValueError('Package phase already ran.')
    bridge = Bridge()
    resolved = {'pip': [], 'npm': []}
    deadline = time.monotonic() + 235
    if packages.get('pip'):
        result = command(['python', '-I', '-m', 'pip', 'install', '--only-binary=:all:', '--no-cache-dir',
                          '--disable-pip-version-check', '--no-input', '--index-url', 'http://127.0.0.1:43123/pypi/simple/',
                          '--trusted-host', '127.0.0.1', '--target', '/deps/python', '--report', '/scratch/pip-report.json',
                          *packages['pip']], max(1, deadline - time.monotonic()), '/scratch')
        if result['status'] != 'complete':
            raise ValueError('pip could not install registry wheels: ' + result['stderr'][-2000:] + (result['error'] or ''))
        report = json.loads(Path('/scratch/pip-report.json').read_text())
        by_name = {re.sub(r'[-_.]+', '-', item['metadata']['name']).lower(): item['metadata'] for item in report['install']}
        for spec in packages['pip']:
            name = re.sub(r'[-_.]+', '-', spec.split('==')[0]).lower()
            info = by_name.get(name)
            if not info:
                raise ValueError('pip did not return a resolved version.')
            resolved['pip'].append(info['name'] + '==' + info['version'])
    if packages.get('npm'):
        result = command(['npm', 'install', '--prefix', '/deps/node', '--ignore-scripts', '--no-audit', '--no-fund',
                          '--save-exact', '--registry', 'http://127.0.0.1:43123/npm/', '--cache', '/scratch/npm-cache',
                          *packages['npm']], max(1, deadline - time.monotonic()), '/scratch')
        if result['status'] != 'complete':
            raise ValueError('npm could not install registry packages without scripts: ' + result['stderr'][-2000:] + (result['error'] or ''))
        data = json.loads(Path('/deps/node/package.json').read_text())
        resolved['npm'] = [name + '@' + version for name, version in data.get('dependencies', {}).items()]
    return {'packages': resolved, 'summary': 'Installed registry wheels and/or npm packages without lifecycle scripts in clean isolated scratch.'}


def files(root=ROOT, count_limit=128):
    result = []
    total = 0
    for directory, dirs, names in os.walk(root, followlinks=False):
        for name in dirs:
            info = os.lstat(Path(directory) / name)
            if not stat.S_ISDIR(info.st_mode):
                raise ValueError('Output contains a link or special directory.')
        for name in names:
            path = Path(directory) / name
            rel = str(path.relative_to(root))
            relative(rel)
            fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
            try:
                info = os.fstat(fd)
                if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_size > FILE_LIMIT:
                    raise ValueError('Output contains a link, special file, or oversized file.')
                total += info.st_size
                if total > TOTAL_LIMIT or len(result) >= count_limit:
                    raise ValueError('Output file budget exceeded.')
                digest = hashlib.sha256()
                while data := os.read(fd, CHUNK):
                    digest.update(data)
                result.append({'path': rel, 'size': info.st_size, 'sha256': digest.hexdigest(), 'mode': 0o755 if info.st_mode & 0o111 else 0o644,
                               'mime': {'.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                        '.pdf': 'application/pdf', '.webm': 'video/webm', '.wav': 'audio/wav',
                                        '.md': 'text/markdown', '.png': 'image/png'}.get(path.suffix.lower(),
                                            mimetypes.guess_type(rel)[0] or 'application/octet-stream')})
            finally:
                os.close(fd)
    return result


def document(path, mime):
    budget = 250000
    passages = []
    if mime == 'application/pdf':
        from pypdf import PdfReader
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ValueError('Encrypted PDFs are not supported.')
        if len(reader.pages) > 200:
            raise ValueError('PDF exceeds 200 pages.')
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            budget -= len(text)
            if budget < 0:
                raise ValueError('Document text exceeds 250000 characters.')
            passages.append({'page': index + 1, 'text': text})
    elif mime == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(item.file_size for item in infos) > 32 * 1024 * 1024:
                raise ValueError('DOCX expanded-size limit exceeded.')
            if any(item.flag_bits & 1 for item in infos):
                raise ValueError('Encrypted DOCX files are unsupported.')
        from docx import Document
        doc = Document(path)
        texts = [paragraph.text for paragraph in doc.paragraphs]
        texts += ['\t'.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        for index, text in enumerate(texts):
            budget -= len(text)
            if budget < 0 or index >= 10000:
                raise ValueError('Document text limit exceeded.')
            passages.append({'paragraph': index + 1, 'text': text})
    else:
        raise ValueError('Unsupported document type.')
    return {'passages': passages}


def probe(path):
    result = command(['ffprobe', '-v', 'error', '-protocol_whitelist', 'file,pipe', '-probesize', '5242880',
                      '-analyzeduration', '5000000', '-show_streams', '-show_format', '-of', 'json', str(path)], 20, cap=512*1024)
    if result['status'] != 'complete':
        raise ValueError('ffprobe could not inspect this file: ' + result['stderr'][-500:])
    data = json.loads(result['stdout'])
    streams = data.get('streams', [])
    audio = next((s for s in streams if s.get('codec_type') == 'audio'), None)
    video = next((s for s in streams if s.get('codec_type') == 'video' and not s.get('disposition', {}).get('attached_pic')), None)
    durations = [data.get('format', {}).get('duration')] + [s.get('duration') for s in streams]
    valid = [float(n) for n in durations if n is not None and n != 'N/A']
    duration = max(valid, default=float('nan'))
    if not math.isfinite(duration) or duration <= 0 or (audio is None and video is None):
        raise ValueError('Media duration or streams are unknown.')
    info = {'durationSeconds': duration, 'hasAudio': audio is not None, 'hasVideo': video is not None}
    if audio:
        info['audioCodec'] = audio.get('codec_name', 'unknown')
    if video:
        info.update(width=int(video.get('width', 0)), height=int(video.get('height', 0)), videoCodec=video.get('codec_name', 'unknown'))
    return info


def handle(message):
    global input_stream, input_hash, input_size, input_expected, sealed
    op = message['op']
    if op == 'hello':
        status = Path('/proc/self/status').read_text()
        return {'uid': os.getuid(), 'pid': os.getpid(), 'noNewPrivileges': 'NoNewPrivs:\t1' in status,
                'seccomp': 'Seccomp:\t2' in status, 'capabilities': next(line for line in status.splitlines() if line.startswith('CapEff:')).split()[1],
                'memory': Path('/sys/fs/cgroup/memory.max').read_text().strip(),
                'swap': Path('/sys/fs/cgroup/memory.swap.max').read_text().strip(),
                'pids': Path('/sys/fs/cgroup/pids.max').read_text().strip(), 'cpu': Path('/sys/fs/cgroup/cpu.max').read_text().strip()}
    if op == 'install':
        return install(message['packages'])
    if op == 'seal':
        if bridge:
            bridge.close()
        sealed = True
        return {}
    if not sealed:
        raise ValueError('Close package access before providing conversation data.')
    if op == 'dependencyManifest':
        return files(Path('/deps'), 4096)
    if op in ('begin', 'beginDependency'):
        if input_stream:
            raise ValueError('An input is already open.')
        path = relative(message['path'])
        if op == 'beginDependency':
            path = Path('/deps') / path.relative_to(ROOT)
        path.parent.mkdir(parents=True, exist_ok=True)
        input_stream = open(path, 'xb')
        os.chmod(path, 0o755 if message.get('mode') == 0o755 else 0o644)
        input_hash, input_size = hashlib.sha256(), 0
        input_expected = message
        return {}
    if op == 'chunk':
        data = base64.b64decode(message['data'], validate=True)
        if not input_stream or len(data) > CHUNK or input_size + len(data) > min(FILE_LIMIT, input_expected['size']):
            raise ValueError('Input chunk exceeds its bound.')
        input_stream.write(data)
        input_hash.update(data)
        input_size += len(data)
        return {}
    if op == 'end':
        input_stream.close()
        input_stream = None
        if input_size != input_expected['size'] or input_hash.hexdigest() != input_expected['sha256']:
            raise ValueError('Input digest mismatch.')
        return {}
    if op == 'execute':
        if message['kind'] not in ('python', 'shell'):
            raise ValueError('Invalid execution kind.')
        argv = ['python', '-c', message['code']] if message['kind'] == 'python' else ['/bin/sh', '-c', message['code']]
        result = command(argv, 60, packages=True)
        result['files'] = files() if result['status'] == 'complete' else []
        return result
    if op == 'document':
        return document(relative(message['path']), message['mime'])
    if op == 'probe':
        return probe(relative(message['path']))
    if op in ('read', 'readDependency'):
        path = relative(message['path'])
        if op == 'readDependency':
            path = Path('/deps') / path.relative_to(ROOT)
        offset = message['offset']
        if not isinstance(offset, int) or offset < 0 or offset > FILE_LIMIT:
            raise ValueError('Invalid file offset.')
        fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
        try:
            info = os.fstat(fd)
            if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_size > FILE_LIMIT:
                raise ValueError('Invalid output file.')
            return {'data': base64.b64encode(os.pread(fd, CHUNK, offset)).decode('ascii')}
        finally:
            os.close(fd)
    raise ValueError('Unknown worker command.')


os.makedirs('/tmp/home', exist_ok=True)
threading.Thread(target=receive, daemon=True).start()
while (request := requests.get()) is not None:
    try:
        result = {'id': request['id'], 'ok': True, 'value': handle(request)}
    except Exception as error:
        result = {'id': request.get('id'), 'ok': False, 'error': str(error)[:2500]}
    emit(result)
